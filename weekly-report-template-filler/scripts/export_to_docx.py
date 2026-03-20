#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def add_markdown_block(doc: Document, block: str) -> None:
    stripped = block.strip()
    if not stripped:
        return
    if stripped.startswith("## "):
        doc.add_heading(stripped[3:].strip(), level=2)
        return
    if stripped.startswith("# "):
        doc.add_heading(stripped[2:].strip(), level=1)
        return
    if stripped.startswith("- "):
        doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        return
    doc.add_paragraph(stripped)


def export_markdown_to_docx(input_path: str, output_path: str) -> None:
    content = Path(input_path).read_text(encoding="utf-8")
    doc = Document()
    for block in content.split("\n\n"):
        lines = [line for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        for line in lines:
            add_markdown_block(doc, line)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export markdown preview to a .docx file.")
    parser.add_argument("--input", required=True, help="Path to markdown file")
    parser.add_argument("--output", required=True, help="Path to .docx output")
    args = parser.parse_args()

    export_markdown_to_docx(args.input, args.output)
    print(f"Wrote docx to {args.output}")


if __name__ == "__main__":
    main()
